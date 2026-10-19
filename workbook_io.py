"""Simple XLSX framework with portable Python imports and exports."""
from datetime import date, datetime, timedelta, time as clock_time
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from pydantic import ValidationError
from domain import Problem
from geocoding import lookup

ROOT = Path(__file__).parent
DAYS = ['Mon','Tue','Wed','Thu','Fri','Sat','Sun']
HEADERS = {
    'Settings':['field','value','help'],
    'Clients':['client_id','name','location_id','address_or_eircode','continuity_cap','preferred_carers','excluded_carers'],
    'Carers':['carer_id','name','skills','max_weekly_hours','unavailable_days','medication_valid_until','dementia_valid_until','colour'],
    'Availability':['carer_id','day','start','end','start_location_id','end_location_id','break_earliest','break_latest','break_minutes'],
    'Visits':['visit_id','client_id','label','days','earliest_start','latest_finish','duration_minutes','skills','fixed_carer_id'],
    'Locations':['location_id','name','address_or_eircode','latitude','longitude'],
    'TravelTimes':['from_location_id','to_location_id','minutes'],
    'Locks':['visit_id','carer_id','start','completed'],
    'Previous':['visit_id','carer_id','start'],
}
REQUIRED = {
    'Settings':['field','value'], 'Clients':['client_id','name'], 'Carers':['carer_id','name','skills'],
    'Availability':['carer_id','day','start','end','start_location_id','end_location_id'],
    'Visits':['visit_id','client_id','days','earliest_start','latest_finish','duration_minutes'],
    'Locations':['location_id','name'], 'TravelTimes':HEADERS['TravelTimes'],
    'Locks':['visit_id','carer_id','start'], 'Previous':HEADERS['Previous'],
}


class WorkbookError(Exception):
    pass


def text(v):
    return '' if v is None else str(v).strip()


def items(v):
    return [s.strip() for s in re.split('[,;]', text(v)) if s.strip()]


def day_list(v):
    raw = text(v).lower().replace('–','-')
    if raw in ['daily','every day']: return list(range(7))
    if raw in ['weekdays','mon-fri']: return list(range(5))
    if raw in ['weekend','weekends']: return [5,6]
    names = {n.lower():i for i,n in enumerate(DAYS)}
    out = []
    for token in items(raw):
        if '-' in token:
            a,b = (names.get(x.strip()[:3]) for x in token.split('-',1))
            if a is None or b is None or a>b: raise ValueError('Use days such as Mon,Tue or Mon-Fri.')
            out.extend(range(a,b+1))
        else:
            if token[:3] not in names: raise ValueError('Use Mon,Tue,Wed,Thu,Fri,Sat,Sun.')
            out.append(names[token[:3]])
    if len(out) != len(set(out)): raise ValueError('A day is listed more than once.')
    return out


def minute(v, *, end=False):
    if isinstance(v,datetime): v=v.time()
    if isinstance(v,timedelta):v=v.total_seconds()/86400
    if isinstance(v,clock_time):
        if v.second or v.microsecond: raise ValueError('Use whole minutes, without seconds.')
        return v.hour*60+v.minute
    if isinstance(v,(int,float)) and not isinstance(v,bool) and 0<=v<=1:
        n=round(v*1440)
        if abs(v*1440-n)>0.001: raise ValueError('Use whole minutes.')
    else:
        match=re.fullmatch(r'(\d{1,2}):(\d{2})(?::00)?',text(v))
        if not match: raise ValueError('Use an Excel time or HH:MM, such as 08:30.')
        h,m=map(int,match.groups())
        if m>59: raise ValueError('Minutes must be 00–59.')
        n=h*60+m
    if not 0<=n<=(1440 if end else 1439): raise ValueError('Time is outside this day.')
    return n


def integer(v, default=None):
    if text(v)=='':
        if default is not None:return default
        raise ValueError('A whole number is required.')
    if isinstance(v,bool): raise ValueError('Use a whole number, not TRUE/FALSE.')
    n=float(v)
    if not n.is_integer():raise ValueError('A whole number is required.')
    return int(n)


def iso_date(v):
    if isinstance(v,(date,datetime)):return v.date().isoformat() if isinstance(v,datetime) else v.isoformat()
    return date.fromisoformat(text(v)).isoformat()


def parse_workbook(content):
    if len(content)>5*1024*1024:raise WorkbookError('Use an .xlsx smaller than 5 MB.')
    try:
        with ZipFile(BytesIO(content)) as z:
            if len(z.infolist())>1000 or sum(x.file_size for x in z.infolist())>25*1024*1024:
                raise WorkbookError('Workbook is too large after decompression.')
            if any('vbaProject' in n or n.startswith('xl/externalLinks/') for n in z.namelist()):
                raise WorkbookError('Use a plain .xlsx without macros or external workbook links.')
        book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
    except (BadZipFile, OSError, KeyError, ValueError) as exc:
        raise WorkbookError('This is not a readable .xlsx workbook. Use the downloadable framework.') from exc
    issues, tables, row_map = [], {}, {}
    def issue(sheet,row,field,message):issues.append(dict(sheet=sheet,row=row,field=field,message=message))
    try:
        for name,expected in HEADERS.items():
            if name not in book.sheetnames:
                if name in ['Settings','Clients','Carers','Availability','Visits']:
                    issue(name,1,'sheet','Required sheet is missing.')
                tables[name]=[];continue
            ws=book[name]
            if ws.max_row is None or ws.max_column is None:
                ws.calculate_dimension(force=True)
            if ws.max_row>1501 or ws.max_column>30:
                issue(name,1,'size','Use at most 1,500 rows and 30 columns per input sheet.');continue
            rows=ws.iter_rows(max_col=3 if name=='Settings' else ws.max_column); header=next(rows,None)
            if not header:issue(name,1,'headers','Header row is missing.');continue
            names=[text(c.value) for c in header]
            if len([n for n in names if n])!=len(set(n for n in names if n)):
                issue(name,1,'headers','Duplicate column names.');continue
            for h in REQUIRED[name]:
                if h not in names:issue(name,1,h,'Required column is missing.')
            for h in names:
                if h and h not in expected:issue(name,1,h,'Unknown input column; keep framework headers unchanged.')
            data=[]
            for cells in rows:
                if not any(c.value is not None and text(c.value)!='' for c in cells):continue
                row=cells[0].row
                if any(c.data_type=='f' for c in cells):
                    issue(name,row,'formula','Input tables accept values only. Paste calculated values before importing.');continue
                data.append((row,dict(zip(names,[c.value for c in cells]))))
            tables[name]=data
    finally:
        book.close()
    if issues:return dict(draft=None,issues=issues,unresolved=[],counts={})
    def build(name,fn):
        out=[]
        for row,record in tables.get(name,[]):
            try:
                value=fn(record,row)
                row_map[(name,len(out))]=row;out.append(value)
            except (ValueError,TypeError,KeyError) as exc:issue(name,row,'row',str(exc))
        return out
    settings={};known={'schema_version','week_start','timezone','continuity_cap','objective','time_limit','travel_buffer'}
    for row,r in tables['Settings']:
        key=text(r.get('field'))
        if key not in known:issue('Settings',row,'field','Unknown setting.');continue
        if key in settings:issue('Settings',row,'field','Duplicate setting.');continue
        settings[key]=r.get('value')
    try:
        if integer(settings.get('schema_version'),1)!=1:raise ValueError('Unsupported schema_version; use 1.')
        settings=dict(week_start=iso_date(settings.get('week_start')),timezone=text(settings.get('timezone')) or 'Europe/Dublin',
            continuity_cap=integer(settings.get('continuity_cap'),3),objective=text(settings.get('objective')) or 'balanced',
            time_limit=integer(settings.get('time_limit'),5),travel_buffer=integer(settings.get('travel_buffer'),0))
    except (ValueError,TypeError) as exc:issue('Settings',2,'value',str(exc))
    locations=build('Locations',lambda r,row:dict(id=text(r.get('location_id')),name=text(r.get('name')),address=text(r.get('address_or_eircode')),
        lat=float(r['latitude']) if text(r.get('latitude')) else None,lng=float(r['longitude']) if text(r.get('longitude')) else None))
    loc_meta={l['id']:dict(sheet='Locations',row=row_map[('Locations',i)],query=l['address']) for i,l in enumerate(locations)}
    def client(r,row):
        cid=text(r.get('client_id')); lid=text(r.get('location_id')); address=text(r.get('address_or_eircode'))
        if not cid or not text(r.get('name')):raise ValueError('client_id and name are required.')
        if not lid:
            if not address:raise ValueError('Provide location_id or address_or_eircode.')
            lid='home_'+cid
            if any(l['id']==lid for l in locations):raise ValueError('Generated home location ID already exists; specify location_id.')
            locations.append(dict(id=lid,name=text(r['name'])+' · home',address=address,lat=None,lng=None))
            loc_meta[lid]=dict(sheet='Clients',row=row,query=address)
        elif address:
            loc=next((l for l in locations if l['id']==lid),None)
            if loc is None:raise ValueError('location_id is not present in Locations.')
            if loc.get('address') and loc['address']!=address:raise ValueError('Address conflicts with this shared location; leave location_id blank to create a new home.')
            loc['address']=address;loc_meta[lid]['query']=address
        return dict(id=cid,name=text(r['name']),location=lid,continuity_cap=integer(r['continuity_cap']) if text(r.get('continuity_cap')) else None,
                    preferred_carers=items(r.get('preferred_carers')),excluded_carers=items(r.get('excluded_carers')))
    clients=build('Clients',client)
    carers=build('Carers',lambda r,row:dict(id=text(r.get('carer_id')),name=text(r.get('name')),skills=items(r.get('skills')),
        max_weekly_minutes=round(float(r.get('max_weekly_hours') if text(r.get('max_weekly_hours')) else 40)*60),
        unavailable=day_list(r.get('unavailable_days')),color=text(r.get('colour')) or '#218477',shifts=[],
        skill_expiry={sk:iso_date(r[col]) for sk,col in [('medication','medication_valid_until'),('dementia','dementia_valid_until')] if text(r.get(col))}))
    availability_rows={}
    for row,r in tables['Availability']:
        try:
            carer=next((c for c in carers if c['id']==text(r.get('carer_id'))),None)
            if carer is None:raise ValueError('carer_id does not match a row in Carers.')
            ds=day_list(r.get('day'))
            if len(ds)!=1:raise ValueError('Use one day per Availability row.')
            s=dict(day=ds[0],start=minute(r.get('start')),end=minute(r.get('end'),end=True),start_location=text(r.get('start_location_id')),end_location=text(r.get('end_location_id')),breaks=[])
            if any(text(r.get(k)) for k in ['break_earliest','break_latest','break_minutes']):
                s['breaks']=[dict(start=minute(r.get('break_earliest')),end=minute(r.get('break_latest'),end=True),duration=integer(r.get('break_minutes')))]
            availability_rows[(carer['id'],len(carer['shifts']))]=row;carer['shifts'].append(s)
        except (ValueError,TypeError) as exc:issue('Availability',row,'row',str(exc))
    visits=build('Visits',lambda r,row:dict(id=text(r.get('visit_id')),client_id=text(r.get('client_id')),label=text(r.get('label')) or 'Care visit',
        days=day_list(r.get('days')),earliest=minute(r.get('earliest_start')),latest=minute(r.get('latest_finish'),end=True),duration=integer(r.get('duration_minutes')),
        skills=items(r.get('skills')) or ['general'],fixed_carer=text(r.get('fixed_carer_id')) or None))
    def lock(r,row):
        completed=text(r.get('completed')).lower()
        if completed and completed not in ['true','false','yes','no','1','0']:raise ValueError('completed must be TRUE or FALSE.')
        return dict(visit_id=text(r.get('visit_id')),carer_id=text(r.get('carer_id')),start=minute(r.get('start')),completed=completed in ['true','yes','1'])
    locks=build('Locks',lock);previous=build('Previous',lock)
    for sheet,records in [('Locations',locations),('Clients',clients),('Carers',carers),('Visits',visits)]:
        seen=set()
        for i,record in enumerate(records):
            row=row_map.get((sheet,i),loc_meta.get(record['id'],{}).get('row',1))
            if record['id'] in seen:issue(sheet,row,'id','Duplicate ID: '+record['id'])
            seen.add(record['id'])
    location_ids={l['id'] for l in locations};carer_ids={c['id'] for c in carers};client_ids={c['id'] for c in clients}
    for i,c in enumerate(clients):
        row=row_map[('Clients',i)]
        if c['location'] not in location_ids:issue('Clients',row,'location_id','Unknown location ID: '+c['location'])
        for cid in c['preferred_carers']+c['excluded_carers']:
            if cid not in carer_ids:issue('Clients',row,'carers','Unknown carer ID: '+cid)
    for i,v in enumerate(visits):
        row=row_map[('Visits',i)]
        if v['client_id'] not in client_ids:issue('Visits',row,'client_id','Unknown client ID: '+v['client_id'])
        if v['fixed_carer'] and v['fixed_carer'] not in carer_ids:issue('Visits',row,'fixed_carer_id','Unknown carer ID.')
    for c in carers:
        for j,s in enumerate(c['shifts']):
            for field in ['start_location','end_location']:
                if s[field] not in location_ids:issue('Availability',availability_rows[(c['id'],j)],field+'_id','Unknown location ID: '+s[field])
    n=len(locations);matrix=[[0 for _ in range(n)] for _ in range(n)]
    if tables.get('TravelTimes'):
        ids=[l['id'] for l in locations];seen=set()
        for row,r in tables['TravelTimes']:
            try:
                a,b=ids.index(text(r.get('from_location_id'))),ids.index(text(r.get('to_location_id')))
                if (a,b) in seen:raise ValueError('Duplicate directed travel-time pair.')
                seen.add((a,b));matrix[a][b]=integer(r.get('minutes'))
            except (ValueError,TypeError) as exc:issue('TravelTimes',row,'row',str(exc))
        if len(seen)!=n*n:issue('TravelTimes',1,'pairs','Provide every from/to pair, including zero diagonals, or remove TravelTimes to calculate automatically.')
    unresolved=[]
    for loc in locations:
        if loc.get('lat') is None and loc.get('lng') is None:
            candidates=lookup(loc['address'],online=False) if len(loc['address'])>=3 else []
            if len(candidates)==1:
                c=candidates[0];loc.update(lat=c['lat'],lng=c['lng'],eircode=c.get('eircode',''),geocode_source=c['provider'],geocode_label=c['label'])
            else:
                unresolved.append(dict(location_id=loc['id'],name=loc['name'],**loc_meta[loc['id']]))
    draft=dict(settings=settings,locations=locations,carers=carers,clients=clients,visits=visits,locks=locks,previous=previous,
               travel_matrix=matrix,routing=dict(kind='pending',provider='Excel import; road times pending',live_traffic=False))
    try:
        draft=Problem.model_validate(draft).model_dump(mode='json')
    except ValidationError as exc:
        sections={'carers':'Carers','clients':'Clients','visits':'Visits','locations':'Locations','locks':'Locks','previous':'Previous','settings':'Settings'}
        for error in exc.errors():
            path=error['loc'];section=sections.get(path[0],'Workbook') if path else 'Workbook';row=1
            if len(path)>1 and isinstance(path[1],int):
                row=row_map.get((section,path[1]),loc_meta.get(locations[path[1]]['id'],{}).get('row',1) if section=='Locations' and path[1]<len(locations) else 1)
                if section=='Carers' and len(path)>3 and path[2]=='shifts':
                    section='Availability';row=availability_rows.get((carers[path[1]]['id'],path[3]),1)
            issue(section,row,' → '.join(map(str,path)),error['msg'])
        draft=None
    counts=dict(clients=len(clients),carers=len(carers),visits=sum(len(v['days']) for v in visits),locations=len(locations),unresolved=len(unresolved))
    return dict(draft=draft,issues=issues,unresolved=unresolved,counts=counts,ready=not issues and not unresolved)


def workbook_spec(p, result=None, *, template=False):
    p=p.model_dump(mode='json') if isinstance(p,Problem) else p
    sheets=[]
    def add(name,headers,rows,**kw):sheets.append(dict(name=name,headers=headers,rows=rows,**kw))
    if result:
        validation=result['validation']; metrics=validation['metrics']; cs={c['id']:c['name'] for c in p['clients']};ks={c['id']:c['name'] for c in p['carers']}
        add('Summary',['measure','value','units'],[['Week starting',p['settings']['week_start'],'local calendar'],['Solver status',result['status'],''],['Verified visits',metrics['visits'],'visits'],['Travel',metrics['travel'],'minutes'],['Avoidable idle',metrics['idle'],'minutes'],['Care',metrics['care'],'minutes'],['Working time',metrics['working'],'minutes'],['Travel source',p.get('routing',{}).get('provider','Manual'),'No live traffic']])
        add('Roster',['date','day','client_id','client','carer_id','carer','visit','start','end','duration_minutes','locked','completed'],
            [[(date.fromisoformat(p['settings']['week_start'])+timedelta(days=v['day'])).isoformat(),DAYS[v['day']],v['client_id'],cs[v['client_id']],r['carer_id'],ks[r['carer_id']],v['label'],v['start']/1440,v['end']/1440,v['duration'],v['locked'],v.get('completed',False)] for r in result['routes'] for v in r['visits']],times=['start','end'],dates=['date'])
        add('Continuity',['client_id','client','carer_ids','distinct_carers','weekly_cap','status'],[[c['id'],c['name'],','.join(validation['continuity'][c['id']]),len(validation['continuity'][c['id']]),c.get('continuity_cap') or p['settings']['continuity_cap'],'Within limit'] for c in p['clients']])
        add('Routes',['carer_id','carer','day','departure','return','break_start','break_end','visit_count'],[[r['carer_id'],ks[r['carer_id']],DAYS[r['day']],r['departure']/1440,r['return']/1440,r['breaks'][0]['start']/1440 if r['breaks'] else None,r['breaks'][0]['end']/1440 if r['breaks'] else None,len(r['visits'])] for r in result['routes']],times=['departure','return','break_start','break_end'])
    help_text=dict(schema_version='Keep 1. Header row is row 1 on every table. Do not rename sheets or headers.',week_start='Monday date, YYYY-MM-DD.',timezone='Use Europe/Dublin.',continuity_cap='Hard weekly limit: 3 or 4 distinct carers per client.',objective='balanced, travel or idle.',time_limit='Search seconds: 1–20.',travel_buffer='Extra minutes per journey between different locations.')
    add('Settings',HEADERS['Settings'],[['schema_version',1,help_text['schema_version']]]+[[k,v,help_text[k]] for k,v in p['settings'].items()],settings=True)
    add('Clients',HEADERS['Clients'],[[c['id'],c['name'],c['location'],'',c.get('continuity_cap'),','.join(c.get('preferred_carers',[])),','.join(c.get('excluded_carers',[]))] for c in p['clients']])
    add('Carers',HEADERS['Carers'],[[c['id'],c['name'],','.join(c['skills']),c['max_weekly_minutes']/60,','.join(DAYS[d] for d in c['unavailable']),c['skill_expiry'].get('medication',''),c['skill_expiry'].get('dementia',''),c['color']] for c in p['carers']],dates=['medication_valid_until','dementia_valid_until'])
    add('Availability',HEADERS['Availability'],[[c['id'],DAYS[s['day']],s['start']/1440,s['end']/1440,s['start_location'],s['end_location'],s['breaks'][0]['start']/1440 if s['breaks'] else None,s['breaks'][0]['end']/1440 if s['breaks'] else None,s['breaks'][0]['duration'] if s['breaks'] else None] for c in p['carers'] for s in c['shifts']],times=['start','end','break_earliest','break_latest'])
    add('Visits',HEADERS['Visits'],[[v['id'],v['client_id'],v['label'],','.join(DAYS[d] for d in v['days']),v['earliest']/1440,v['latest']/1440,v['duration'],','.join(v['skills']),v.get('fixed_carer') or ''] for v in p['visits']],times=['earliest_start','latest_finish'])
    def address_label(l):
        value=l.get('address') or l['name'];code=l.get('eircode','')
        return value+', '+code if code and code.replace(' ','').upper() not in value.replace(' ','').upper() else value
    add('Locations',HEADERS['Locations'],[[l['id'],l['name'],address_label(l),l.get('lat'),l.get('lng')] for l in p['locations']])
    if not template:
        if p.get('routing',{}).get('kind')!='pending':
            add('TravelTimes',HEADERS['TravelTimes'],[[a['id'],b['id'],p['travel_matrix'][i][j]] for i,a in enumerate(p['locations']) for j,b in enumerate(p['locations'])])
        for sheet,attr in [('Locks','locks'),('Previous','previous')]:
            if p.get(attr):add(sheet,HEADERS[sheet],[[l['visit_id'],l['carer_id'],l['start']/1440]+([l.get('completed',False)] if sheet=='Locks' else []) for l in p[attr]],times=['start'])
    return dict(sheets=sheets,notes=[
        'Edit the six input sheets. Add rows directly below the tables. Keep IDs unique and consistent across sheets.',
        'Clients: use location_id from Locations OR leave it blank and enter address_or_eircode to create a home location.',
        'Locations: enter an address/Eircode; coordinates are optional. Carer start/end locations must reference a Locations ID.',
        'Days: Mon,Tue,Wed,Thu,Fri,Sat,Sun; daily and Mon-Fri are also accepted. Availability uses one day per row.',
        'Skills: general,medication,dementia, comma-separated. Blank Visits skills defaults to general.',
        'Use Excel times or HH:MM. Latest finish includes the full visit duration. Weekly hours include care, travel and waiting; breaks excluded.',
        'Input cells must contain values, not formulas. Clear example rows before entering your own data.',
        'Demo lookup: public places only. Try Trinity College Dublin / D02 PN40, Dublin Castle / D02 XN27, Guinness Storehouse / D08 VF8H.',
        'No-provider lookup is not a complete Eircode database. Review candidates; unresolved locations prevent applying the import.',
        'Import preview does not replace the current draft. Apply only after resolving issues; road times are calculated automatically.',
        'Output tabs are snapshots. Reimport input sheets and regenerate in the app after edits; Excel does not run OR-Tools.',
        'Existing locks/completed visits are preserved in Locks when exporting the current scenario. A new template starts a fresh week.',
        'Public demo Eircode source, Trinity College: https://www.tcd.ie/contacts/',
        'Public demo Eircode source, Dublin Castle: https://heritageireland.ie/visit/places-to-visit/dublin-castle/',
        'Public demo Eircode source, Guinness Storehouse: https://www.guinness-storehouse.com/en/contact-us',
    ])


def export_workbook(p,result=None,*,template=False,output_path=None,render=False,render_sheets=None):
    from xlsx_writer import write_xlsx
    if render:
        raise WorkbookError('Image rendering is not part of the portable Excel exporter.')
    body=write_xlsx(workbook_spec(p,result,template=template))
    if output_path:
        target=Path(output_path);target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(body)
    return body
