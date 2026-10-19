from datetime import timedelta
from io import BytesIO
import json
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
import app as server
from domain import Problem
from examples import examples
from geocoding import lookup, eircode_in
from scheduler import solve
from workbook_io import parse_workbook, export_workbook, workbook_spec, minute, day_list

TEMPLATE=Path(__file__).resolve().parents[1]/'templates/HomeCare-template.xlsx'
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def fixture_edit(sheet_number,changes):
    """Mutate fixture XML only, for negative/extension importer tests."""
    source=TEMPLATE.read_bytes();out=BytesIO();path=f'xl/worksheets/sheet{sheet_number}.xml'
    with ZipFile(BytesIO(source)) as original,ZipFile(out,'w',ZIP_DEFLATED) as target:
        for item in original.infolist():
            data=original.read(item.filename)
            if item.filename==path:
                root=ET.fromstring(data);body=root.find('s:sheetData',NS)
                for addr,value in changes.items():
                    cell=root.find(f'.//s:c[@r="{addr}"]',NS)
                    if cell is None:
                        row_number=''.join(x for x in addr if x.isdigit());row=root.find(f'.//s:row[@r="{row_number}"]',NS)
                        if row is None:row=ET.SubElement(body,'{'+NS['s']+'}row',r=row_number)
                        cell=ET.SubElement(row,'{'+NS['s']+'}c',r=addr)
                    cell.clear();cell.set('r',addr)
                    if value is not None:
                        cell.set('t','inlineStr');node=ET.SubElement(cell,'{'+NS['s']+'}is');ET.SubElement(node,'{'+NS['s']+'}t').text=value
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            target.writestr(item,data)
    return out.getvalue()


def test_template_is_six_sheets_and_imports_every_visit():
    b=load_workbook(TEMPLATE,read_only=True);assert b.sheetnames==['Settings','Clients','Carers','Availability','Visits','Locations'];b.close()
    r=parse_workbook(TEMPLATE.read_bytes())
    assert not r['issues'] and r['ready'] and r['counts']['visits']==56
    assert r['draft']['carers'][0]['shifts'][0]['start']==480
    assert r['draft']['settings']['week_start']=='2026-10-19'


def test_direct_client_eircode_creates_home_without_coordinates():
    r=parse_workbook(fixture_edit(2,{'C2':None,'D2':'D02PN40'}))
    assert not r['issues'] and not r['unresolved']
    c=r['draft']['clients'][0];l=next(l for l in r['draft']['locations'] if l['id']==c['location'])
    assert l['eircode']=='D02 PN40' and l['lat']==53.3438
    assert r['counts']['visits']==56


def test_unknown_address_stays_unresolved_not_guessed():
    r=parse_workbook(fixture_edit(2,{'C2':None,'D2':'Unknown public demo place 999'}))
    assert len(r['unresolved'])==1 and not r['ready']
    assert r['unresolved'][0]['sheet']=='Clients' and r['unresolved'][0]['row']==2


def test_unknown_client_reference_reports_exact_excel_row():
    r=parse_workbook(fixture_edit(5,{'B2':'no_such_client'}))
    assert any(i['sheet']=='Visits' and i['row']==2 and i['field']=='client_id' for i in r['issues'])


def test_duplicate_client_is_not_silently_accepted():
    r=parse_workbook(fixture_edit(2,{'A3':'mary'}))
    assert any(i['sheet']=='Clients' and i['row']==3 and 'Duplicate' in i['message'] for i in r['issues'])


def test_added_input_row_is_read_and_preserved():
    r=parse_workbook(fixture_edit(5,{'A10':'extra','B10':'mary','C10':'Extra care','D10':'Mon','E10':'14:00','F10':'14:45','G10':'30','H10':'general'}))
    assert not r['issues'] and r['counts']['visits']==57


def test_bad_time_and_invalid_whole_window_are_visible():
    r=parse_workbook(fixture_edit(5,{'E2':'banana'}));assert any(i['sheet']=='Visits' and i['row']==2 for i in r['issues'])
    r=parse_workbook(fixture_edit(5,{'G2':'240'}));assert any('entire visit' in i['message'] for i in r['issues'])


def test_time_end_of_day_and_day_ranges():
    assert minute(timedelta(days=1),end=True)==1440
    assert minute('24:00',end=True)==1440
    assert day_list('Mon-Fri')==[0,1,2,3,4]
    with pytest.raises(ValueError):day_list('Mon,Mon')


def test_known_public_eircodes_and_no_online_calls(monkeypatch):
    monkeypatch.setattr('geocoding.urlopen',lambda *a,**k:pytest.fail('Should use local demo gazetteer'))
    assert lookup('d02pn40',online=True)[0]['name']=='Trinity College Dublin'
    assert lookup('D02 XN27',online=True)[0]['name']=='Dublin Castle'
    assert eircode_in('D6W ABC1')=='D6WABC1'


def test_public_name_does_not_override_a_conflicting_eircode():
    assert lookup('Trinity College Dublin D02 ZZZ9',online=False)==[]


def test_pending_export_does_not_label_placeholder_times_as_estimates():
    p=examples()[0]['data']
    p['routing']={'kind':'pending','provider':'Pending','live_traffic':False}
    p['locations'][0].update(address='Trinity College Dublin',eircode='D02 PN40')
    spec=workbook_spec(p)
    assert 'TravelTimes' not in [s['name'] for s in spec['sheets']]
    assert next(s for s in spec['sheets'] if s['name']=='Locations')['rows'][0][2]=='Trinity College Dublin, D02 PN40'


def test_excel_results_export_rejects_pending_travel_and_malformed_result():
    client=TestClient(server.app);p=examples()[0]['data']
    p['routing']={'kind':'pending','provider':'Pending','live_traffic':False}
    assert client.post('/api/excel/export',json={'problem':p,'result':{'routes':[],'status':'FEASIBLE'}}).status_code==422
    assert client.post('/api/excel/export',json={'problem':p,'result':[]}).status_code==422


def test_eircode_query_never_degrades_to_postal_district(monkeypatch,tmp_path):
    class Reply:
        def __enter__(self):return BytesIO(json.dumps({'features':[{'properties':{'name':'Dublin 2','postcode':'D02','osm_value':'suburb'},'geometry':{'coordinates':[-6.26,53.34]}}]}).encode())
        def __exit__(self,*args):pass
    monkeypatch.setattr('geocoding.CACHE',tmp_path);monkeypatch.setattr('geocoding.urlopen',lambda *args,**kwargs:Reply());monkeypatch.setattr('geocoding._last',0)
    assert lookup('D02 ZZZ9',online=True)==[]


def test_excel_api_preview_does_not_apply_or_solve(monkeypatch):
    client=TestClient(server.app)
    monkeypatch.setattr(server,'solve',lambda *a,**k:pytest.fail('Import must not start solving'))
    r=client.post('/api/excel/import',files={'file':('demo.xlsx',TEMPLATE.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert r.status_code==200 and r.json()['counts']['visits']==56
    assert client.get('/api/excel/template').status_code==200
    assert client.post('/api/excel/import',files={'file':('bad.xls',b'bad')}).status_code==422
    assert client.post('/api/geocode',json={'query':'D02 PN40','public_demo':True}).json()['candidates'][0]['lat']==53.3438


def test_export_results_round_trip_preserves_inputs_and_locks():
    p=Problem.model_validate(examples()[1]['data']);r=solve(p)
    data=p.model_dump(mode='json');v=r['routes'][0]['visits'][0]
    data['locks']=[dict(visit_id=v['id'],carer_id=v['carer_id'],start=v['start'],completed=True)]
    p=Problem.model_validate(data)
    exported=export_workbook(p,r)
    imported=parse_workbook(exported)
    assert not imported['issues'] and imported['counts']['visits']==5
    assert imported['draft']['locks'][0]['completed'] is True
    assert imported['draft']['travel_matrix']==data['travel_matrix']
    book=load_workbook(BytesIO(exported),read_only=True)
    assert {'Summary','Roster','Routes','Continuity'}<=set(book.sheetnames)
    assert sum(1 for row in book['Roster'].iter_rows(min_row=2,values_only=True) if row[0])==5
    book.close()
