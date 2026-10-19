"""Verify the packaged HTTP service without external map or lookup calls."""
import sys
from io import BytesIO
import httpx
from openpyxl import load_workbook

with httpx.Client(base_url=sys.argv[1],timeout=60) as client:
    response=client.get('/api/health');response.raise_for_status()
    assert response.json()['scheduler_revision']=='scalability-v3'
    assert client.get('/').status_code==200
    assert client.get('/static/app.js').status_code==200
    template=client.get('/api/excel/template');template.raise_for_status()
    preview=client.post('/api/excel/import',files={'file':('template.xlsx',template.content)})
    preview.raise_for_status();assert preview.json()['counts']['visits']==56
    problem=client.get('/api/examples').json()[1]['data']
    solved=client.post('/api/solve',json=problem);solved.raise_for_status();result=solved.json()
    assert result['status'] in ['FEASIBLE','OPTIMAL'] and result['validation']['valid']
    exported=client.post('/api/excel/export',json={'problem':problem,'result':result})
    exported.raise_for_status();book=load_workbook(BytesIO(exported.content),read_only=True)
    assert {'Roster','Routes','Continuity','Summary'}<=set(book.sheetnames);book.close()
    print('HTTP deployment smoke passed: UI, template/import, solver and Excel result export.')
