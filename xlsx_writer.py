"""Portable Excel serialization using only the project's Python dependencies."""
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


def write_xlsx(spec):
    book=Workbook();book.remove(book.active)
    for table in spec['sheets']:
        sheet=book.create_sheet(table['name']);sheet.sheet_view.showGridLines=False
        headers=table['headers'];sheet.append(headers)
        for record in table['rows']:
            sheet.append(record)
        if not table['rows']:sheet.append([None]*len(headers))
        for row in sheet:
            for cell in row:
                # Names and addresses are literal values, including strings
                # beginning with '='. Never convert user input to formulas.
                if isinstance(cell.value,str):cell.data_type='s'
                cell.font=Font(name='Arial',size=10,color='243B39')
                cell.alignment=Alignment(vertical='center')
            sheet.row_dimensions[row[0].row].height=25
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='217467')
            cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[1].height=38;sheet.freeze_panes='A2'
        for index,name in enumerate(headers,1):
            column=get_column_letter(index);sheet.column_dimensions[column].width=22
            if name in ['name','client','carer','label','address_or_eircode','visit']:
                sheet.column_dimensions[column].width=32
            for row in range(2,sheet.max_row+1):
                cell=sheet.cell(row,index)
                if name in table.get('times',[]):
                    cell.number_format='[h]:mm' if name in ['end','latest_finish','return','break_latest','break_end'] else 'hh:mm'
                if name in table.get('dates',[]) and cell.value:
                    cell.value=date.fromisoformat(str(cell.value));cell.number_format='yyyy-mm-dd'
                if name in ['latitude','longitude']:cell.number_format='0.000000'
                if name=='max_weekly_hours':cell.number_format='0.00'
                if name in ['address_or_eircode','skills','days','preferred_carers','excluded_carers','unavailable_days']:
                    cell.alignment=Alignment(wrap_text=True,vertical='center');sheet.row_dimensions[row].height=39
            if name in ['continuity_cap','day']:
                values='3,4' if name=='continuity_cap' else 'Mon,Tue,Wed,Thu,Fri,Sat,Sun'
                validation=DataValidation(type='list',formula1='"'+values+'"',allow_blank=True)
                validation.showErrorMessage=True;sheet.add_data_validation(validation);validation.add(f'{column}2:{column}1501')
        excel_table=Table(displayName=table['name']+'Table',ref=f'A1:{get_column_letter(len(headers))}{sheet.max_row}')
        excel_table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        sheet.add_table(excel_table)
        if table.get('settings'):
            sheet.column_dimensions['C'].width=72
            for row in range(2,sheet.max_row+1):
                sheet.cell(row,3).alignment=Alignment(wrap_text=True,vertical='center');sheet.row_dimensions[row].height=36
                if sheet.cell(row,1).value=='week_start':
                    sheet.cell(row,2).value=date.fromisoformat(sheet.cell(row,2).value);sheet.cell(row,2).number_format='yyyy-mm-dd'
            sheet.merge_cells('E1:J1');sheet['E1']='How to use this workbook'
            sheet['E1'].font=Font(name='Arial',size=14,bold=True,color='243B39')
            for index,note in enumerate(spec['notes'],3):
                sheet.merge_cells(start_row=index,start_column=5,end_row=index,end_column=10)
                cell=sheet.cell(index,5,f'{index-2}. {note}');cell.data_type='s'
                cell.font=Font(name='Arial',size=10,color='52695B');cell.alignment=Alignment(wrap_text=True,vertical='center')
                sheet.row_dimensions[index].height=40
            for column in range(5,11):sheet.column_dimensions[get_column_letter(column)].width=14
    stream=BytesIO();book.save(stream);book.close();return stream.getvalue()
